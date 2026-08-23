from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parent
REP=ROOT/'report'; REP.mkdir(exist_ok=True)
df=pd.read_csv(ROOT/'data'/'cleaned_nassau_candy.csv',parse_dates=['Order Date','Ship Date'])
p=pd.read_csv(ROOT/'outputs'/'product_profitability.csv')
d=pd.read_csv(ROOT/'outputs'/'division_performance.csv')

sales=df.Sales.sum(); profit=df['Gross Profit'].sum(); cost=df.Cost.sum(); units=df.Units.sum(); gm=profit/sales
ch=d.loc[d.Division=='Chocolate'].iloc[0]; other=d.loc[d.Division=='Other'].iloc[0]; sugar=d.loc[d.Division=='Sugar'].iloc[0]
kaz=p[p['Product Name'].eq('Kazookles')].iloc[0]
lic=p[p['Product Name'].eq('Lickable Wallpaper')].iloc[0]
top_profit=p.iloc[0]
rev80=(p.sort_values('Sales',ascending=False).Sales.cumsum()/p.Sales.sum()<.8).sum()+1
prof80=(p.sort_values('Gross_Profit',ascending=False).Gross_Profit.cumsum()/p.Gross_Profit.sum()<.8).sum()+1

# ---------- helpers ----------
def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run(); fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text='PAGE'; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); run._r.extend([fldChar1,instr,fldChar2])

def style_doc(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5)
    for s,size,color in [('Title',26,'1F4E79'),('Heading 1',18,'1F4E79'),('Heading 2',13,'2F5597'),('Heading 3',11,'404040')]:
        styles[s].font.name='Aptos Display'; styles[s].font.size=Pt(size); styles[s].font.color.rgb=RGBColor.from_string(color)
    footer=sec.footer.paragraphs[0]; add_page_number(footer)

def add_title(doc,title,subtitle=None):
    p=doc.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    if subtitle:
        q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run(subtitle); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(90,90,90)

def add_kpi_table(doc,items):
    t=doc.add_table(rows=2,cols=len(items)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,(label,val) in enumerate(items):
        c=t.cell(0,i); c.text=label; set_cell_shading(c,'D9EAF7'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c2=t.cell(1,i); c2.text=val; c2.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for rr in [c.paragraphs[0],c2.paragraphs[0]]: rr.alignment=WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].bold=True; c2.paragraphs[0].runs[0].bold=True; c2.paragraphs[0].runs[0].font.size=Pt(12)

def add_bullets(doc,arr):
    for x in arr: doc.add_paragraph(x,style='List Bullet')

def add_img(doc,name,width=6.8):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ROOT/'assets'/name),width=Inches(width))

def add_product_table(doc):
    cols=['Product Name','Division','Sales','Gross_Profit','Gross_Margin','Profit_per_Unit','Margin_Risk_Flag']
    t=doc.add_table(rows=1,cols=len(cols)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    labels=['Product','Division','Sales','Gross Profit','Margin','Profit/Unit','Risk']
    for i,l in enumerate(labels): t.rows[0].cells[i].text=l; set_cell_shading(t.rows[0].cells[i],'1F4E79'); t.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255); t.rows[0].cells[i].paragraphs[0].runs[0].bold=True
    for _,r in p.iterrows():
        row=t.add_row().cells
        vals=[r['Product Name'],r.Division,f"${r.Sales:,.2f}",f"${r.Gross_Profit:,.2f}",f"{r.Gross_Margin:.1%}",f"${r.Profit_per_Unit:,.2f}",r.Margin_Risk_Flag]
        for i,v in enumerate(vals): row[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0); par.paragraph_format.space_before=Pt(0)
                for run in par.runs: run.font.size=Pt(7.3)

# ---------- research report ----------
doc=Document(); style_doc(doc)
add_title(doc,'Product Line Profitability & Margin Performance Analysis','Nassau Candy Distributor | End-to-End Analytics Project')
p0=doc.add_paragraph(); p0.alignment=WD_ALIGN_PARAGRAPH.CENTER; p0.add_run('Dataset: 10,194 order-line records | Order period: 2024-01-02 to 2025-12-31').italic=True
add_kpi_table(doc,[('Sales',f'${sales:,.0f}'),('Gross Profit',f'${profit:,.0f}'),('Gross Margin',f'{gm:.1%}'),('Units',f'{units:,.0f}'),('Products',f'{p.shape[0]}')])

doc.add_heading('Executive Summary',level=1)
doc.add_paragraph(f'Nassau Candy Distributor generated ${sales:,.2f} in sales and ${profit:,.2f} in gross profit, producing an overall gross margin of {gm:.1%}. Profitability is strong at the portfolio level, but it is highly concentrated: the Chocolate division contributes {ch.Revenue_Share:.1%} of revenue and {ch.Profit_Share:.1%} of gross profit. The five major chocolate products alone account for at least 80% of both revenue and profit.')
add_bullets(doc,[
    f"The leading product by gross profit is {top_profit['Product Name']} at ${top_profit.Gross_Profit:,.2f}, with a {top_profit.Gross_Margin:.1%} margin.",
    f"Kazookles is the clearest margin-risk item: only {kaz.Gross_Margin:.1%} gross margin and ${kaz.Profit_per_Unit:.2f} profit per unit; cost absorbs {(kaz.Cost/kaz.Sales):.1%} of sales.",
    f"Lickable Wallpaper produces a relatively high ${lic.Profit_per_Unit:.2f} profit per unit but only a {lic.Gross_Margin:.1%} margin, creating a different pricing/cost-management profile.",
    'There is no evidence of a truly high-sales / low-margin product among the dominant chocolate lines; all five major chocolate products maintain margins above roughly 64%.',
    'Margin volatility is effectively zero at product level because the dataset uses stable product economics over time; volatility should therefore not be treated as a differentiating KPI until richer price/cost history is available.'
])

doc.add_heading('1. Business Problem and Objectives',level=1)
doc.add_paragraph('Sales volume can overstate product health when cost-to-serve or manufacturing cost consumes most of the selling price. This project converts transaction-level order data into product and division profitability intelligence to support pricing, sourcing, promotion and assortment decisions.')
add_bullets(doc,['Identify products that create the most absolute gross profit.','Separate margin quality from sales volume.','Compare division-level economic efficiency and concentration.','Quantify Pareto dependency for revenue and profit.','Flag products for repricing, cost renegotiation or discontinuation review.','Deliver a Streamlit application that makes the analysis filterable and repeatable.'])

doc.add_heading('2. Data Cleaning and Validation',level=1)
doc.add_paragraph('The source contains 10,194 rows and 18 original fields. No missing values were present in Sales, Cost, Gross Profit or Units. No duplicate rows or duplicate Row IDs were found. Sales and Units were positive in all records, and the accounting identity Sales - Cost = Gross Profit reconciled to floating-point precision for every row.')
add_bullets(doc,[
    'Order Date and Ship Date were parsed with day-first date logic.',
    'Product/division labels were trimmed and whitespace standardized; the Scrumdiddlyumptious naming variant was normalized.',
    'Factory, latitude and longitude were added using the supplied product-to-factory mapping.',
    'Gross Margin = Gross Profit / Sales; Profit per Unit = Gross Profit / Units.',
    'Revenue and Profit Contribution are calculated against filtered portfolio totals.',
    'Invalid rows would be excluded where Sales <= 0, Units <= 0 or required financial fields are null; none were excluded in the supplied file.'
])
pwarn=doc.add_paragraph(); rw=pwarn.add_run('Data-quality caveat: '); rw.bold=True; pwarn.add_run('Ship dates range from 2026-06-30 to 2030-06-28 while order dates range from 2024-01-02 to 2025-12-31. Therefore this profitability project does not infer delivery duration, shipping delay or route efficiency from those fields.')

doc.add_heading('3. Portfolio-Level Financial Performance',level=1)
add_kpi_table(doc,[('Sales',f'${sales:,.2f}'),('Cost',f'${cost:,.2f}'),('Gross Profit',f'${profit:,.2f}'),('Margin',f'{gm:.2%}'),('Profit / Unit',f'${profit/units:.2f}')])
doc.add_paragraph('At aggregate level, the portfolio converts nearly two-thirds of sales into gross profit. This is a healthy headline margin, but aggregate performance masks meaningful variation across products and divisions.')
add_img(doc,'product_gross_profit.png')

doc.add_heading('4. Product-Level Profitability',level=1)
doc.add_paragraph('The five chocolate bars dominate absolute profit. Their margins range from about 64.9% to 71.3%, demonstrating that the highest-volume products are not sacrificing margin materially. Product risk instead resides in smaller non-core items where cost intensity is high.')
add_product_table(doc)

doc.add_heading('5. Division-Level Performance',level=1)
add_img(doc,'division_revenue_profit.png',6.6)
doc.add_paragraph(f'Chocolate is the economic engine: ${ch.Sales:,.2f} sales, ${ch.Gross_Profit:,.2f} gross profit and {ch.Gross_Margin:.1%} margin. Other contributes {other.Revenue_Share:.1%} of revenue but only {other.Profit_Share:.1%} of profit, with a lower {other.Gross_Margin:.1%} margin. Sugar is very small ({sugar.Revenue_Share:.2%} of revenue) but its {sugar.Gross_Margin:.1%} margin is broadly healthy, meaning its issue is scale rather than margin quality.')

doc.add_heading('6. Cost Structure and Margin Diagnostics',level=1)
add_img(doc,'sales_vs_margin.png',6.6)
doc.add_paragraph(f'Kazookles is the primary structural margin problem: cost is ${kaz.Cost:,.2f} against ${kaz.Sales:,.2f} sales, leaving only ${kaz.Gross_Profit:,.2f} gross profit. This equates to a {(kaz.Cost/kaz.Sales):.1%} cost ratio and {kaz.Gross_Margin:.1%} margin. By contrast, Lickable Wallpaper has a 50.0% margin and a high profit per unit; its economics justify a pricing/sourcing review rather than automatic discontinuation.')
add_bullets(doc,[
    'Critical (<20% margin): urgent reprice, cost renegotiation or discontinuation review.',
    'High risk (20%-50%): validate market price ceiling and supplier/manufacturing cost.',
    'Watch (50%-60%): monitor mix and test pricing where demand is resilient.',
    'Healthy (>60%): protect availability and scale selectively, while watching concentration risk.'
])

doc.add_heading('7. Profit Concentration (Pareto) Analysis',level=1)
add_img(doc,'pareto_profit.png',6.7)
doc.add_paragraph(f'{prof80} of {len(p)} products ({prof80/len(p):.1%}) are sufficient to cross 80% of gross profit. The same count ({rev80} products) crosses 80% of revenue. This is a strong focus advantage for assortment management, but it creates dependency risk: disruption, price pressure or supply issues affecting the core chocolate products would disproportionately impact enterprise results.')

doc.add_heading('8. Margin Volatility',level=1)
doc.add_paragraph('Monthly product margin standard deviation is approximately zero for all products. This is not evidence that real-world margins are inherently stable; it reflects the synthetic/structured nature of the supplied economics, where product price-cost relationships remain nearly constant. A production-grade model should incorporate supplier cost changes, promotional discounts, freight allocation and actual transaction price movements.')

doc.add_heading('9. Recommendations',level=1)
recs=[
('1. Put Kazookles into an immediate margin recovery workstream','Model price increases, packaging/recipe changes, factory cost renegotiation and minimum-order economics. Establish a target margin floor before deciding whether to retain the SKU.'),
('2. Review Lickable Wallpaper for price architecture','Its absolute profit per unit is attractive but the 50% margin trails the core chocolate lines. Test whether price can move upward without materially reducing volume.'),
('3. Protect the five core chocolate SKUs','They are the principal profit pool. Prioritize supply continuity, service levels, selective promotions and inventory availability; avoid promotions that dilute margin without incremental volume.'),
('4. Manage concentration explicitly','Track profit contribution by top 5 products and by factory. Set contingency sourcing and operational plans because core products are concentrated in a small number of factories.'),
('5. Treat Sugar as a growth experiment, not a margin repair problem','Sugar margins are generally acceptable but sales contribution is tiny. Test distribution, placement and promotions with ROI gates rather than broad cost-cutting.'),
('6. Improve cost accounting and volatility data','Add freight, warehousing, rebates, discounts, supplier price changes and returns to move from gross-margin analysis toward true contribution margin.'),
('7. Operationalize the dashboard','Review threshold flags monthly. Make SKU action status (maintain/reprice/renegotiate/discontinue) part of the commercial operating cadence.')]
for h,b in recs:
    doc.add_heading(h,level=2); doc.add_paragraph(b)

doc.add_heading('10. Streamlit Dashboard Design',level=1)
add_bullets(doc,['Overview: product gross-profit ranking, margin leaderboard and threshold risk table.','Division Performance: revenue vs profit and division margins.','Cost & Margin Diagnostics: cost-sales and sales-margin scatter plots plus recommended actions.','Pareto Concentration: switchable Sales/Gross Profit cumulative analysis and 80% dependency indicator.','Data Explorer: filtered transaction-level table and CSV download.','Controls: order-date range, division multiselect, margin-threshold slider and product search.'])

doc.add_heading('11. Limitations and Next Data Needed',level=1)
add_bullets(doc,['Gross profit does not include all distribution and cost-to-serve elements; contribution margin would be more decision-relevant.','No promotion/discount field is available, so price elasticity and promo profitability cannot be assessed.','Ship dates are not reliable for lead-time analysis in this file.','Margin volatility is not informative because product economics are nearly fixed.','Customer-level profitability requires allocation of service, freight and handling costs.'])

doc.add_heading('12. Conclusion',level=1)
doc.add_paragraph('The analysis establishes a clear, data-driven view of product-line profitability for Nassau Candy Distributor. The portfolio is financially strong overall, but profit is concentrated in five chocolate products. The priority is not broad-based margin repair; it is targeted action on margin-poor SKUs, protection of the core chocolate profit pool, and deliberate management of product/factory concentration. The accompanying Streamlit dashboard converts these findings into a repeatable management tool for pricing, sourcing and portfolio decisions.')

rp=REP/'Nassau_Candy_Profitability_Research_Report.docx'; doc.save(rp)

# ---------- executive summary ----------
ex=Document(); style_doc(ex)
add_title(ex,'Executive Stakeholder Summary','Nassau Candy Distributor - Product Line Profitability & Margin Performance')
add_kpi_table(ex,[('Sales',f'${sales:,.0f}'),('Gross Profit',f'${profit:,.0f}'),('Gross Margin',f'{gm:.1%}'),('Profit / Unit',f'${profit/units:.2f}')])
ex.add_heading('Decision Message',level=1)
ex.add_paragraph(f'The portfolio is profitable overall, but financial performance is highly concentrated. Chocolate produces {ch.Profit_Share:.1%} of gross profit, and only five of fifteen products are needed to exceed 80% of total profit. Management should protect the core chocolate profit pool while using targeted corrective actions on margin-poor peripheral products.')
ex.add_heading('Priority Findings',level=1)
add_bullets(ex,[
    f"Top profit product: {top_profit['Product Name']} - ${top_profit.Gross_Profit:,.0f} gross profit at {top_profit.Gross_Margin:.1%} margin.",
    f"Primary margin risk: Kazookles - {kaz.Gross_Margin:.1%} margin and only ${kaz.Profit_per_Unit:.2f} profit per unit.",
    f"Other division efficiency gap: {other.Revenue_Share:.1%} of revenue versus {other.Profit_Share:.1%} of profit.",
    f"Sugar division is small ({sugar.Revenue_Share:.2%} of sales) but not structurally margin-poor ({sugar.Gross_Margin:.1%} margin).",
    'No major high-sales/low-margin SKU was found among the five dominant chocolate products.'
])
ex.add_heading('Recommended Actions',level=1)
add_bullets(ex,['Urgently review Kazookles pricing, cost base and portfolio fit.','Test price/cost improvement for Lickable Wallpaper before considering rationalization.','Protect supply and commercial availability for the five core chocolate products.','Add contribution-margin inputs such as freight, warehousing, discounts and returns.','Use the Streamlit dashboard as a monthly SKU margin governance tool.'])
ex.add_heading('Governance / Data Caveat',level=1)
ex.add_paragraph('The supplied Ship Date values are inconsistent with the Order Date period, so shipping-duration or route-performance conclusions are excluded. The analysis is restricted to validated financial fields and product/division structures.')
expath=REP/'Nassau_Candy_Executive_Stakeholder_Summary.docx'; ex.save(expath)
print(rp); print(expath)
